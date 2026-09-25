from __future__ import annotations

import json
import re
import subprocess
import zipfile
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
MANUSCRIPT = ROOT / "manuscript"
SUPPLEMENT = ROOT / "supplement"
TABLES = ROOT / "tables"
FIGURES = ROOT / "figures"
RESULTS = ROOT / "results"
METADATA = pd.read_csv(ROOT / "analysis" / "reference_metadata.csv")
VALUES = json.loads((RESULTS / "manuscript_values.json").read_text())

REFERENCE_ORDER = [
    "10.1093/qje/qjag026",
    "10.5751/es-12875-270209",
    "10.14989/68028",
    "10.4102/hts.v69i1.1175",
    "10.1093/ije/dyh299",
    "10.1126/science.185.4157.1124",
    "10.1037/h0055873",
    "10.1016/j.ecolecon.2011.11.011",
    "10.1016/j.ecolecon.2015.01.006",
    "10.1016/j.ecolecon.2013.07.002",
    "10.1016/j.cosust.2016.12.006",
    "10.1016/j.ecoser.2016.11.007",
    "10.1016/j.cosust.2018.09.005",
    "10.1016/j.ecolecon.2024.108425",
    "10.1029/2000wr900320",
    "10.1111/j.1475-4932.2006.00293.x",
    "10.1086/261849",
    "10.1016/0022-0531(82)90030-8",
    "10.1002/qj.3803",
    "10.1111/j.2517-6161.1972.tb00899.x",
]

FIGURE_FILES = {
    "FIGURE_1": FIGURES / "figure1_framework.png",
    "FIGURE_2": FIGURES / "figure2_empirical_hazard.png",
    "FIGURE_3": FIGURES / "figure3_pseudo_event_success.png",
    "FIGURE_4": FIGURES / "figure4_simulation_decomposition.png",
    "FIGURE_5": FIGURES / "figure5_robustness.png",
}

TABLE_FILES = {
    "TABLE_1": TABLES / "table1_study_locations.csv",
    "TABLE_2": TABLES / "table2_pseudo_event_results.csv",
    "TABLE_3": TABLES / "table3_simulation_results.csv",
    "TABLE_4": TABLES / "table4_hypothesis_classification.csv",
    "TABLE_S1": TABLES / "table_s1_price_models.csv",
    "TABLE_S2": TABLES / "table_s2_threshold_window_sensitivity.csv",
    "TABLE_S3": TABLES / "table_s3_parameter_sensitivity.csv",
    "TABLE_S4": TABLES / "table_s4_simulation_convergence.csv",
    "TABLE_S5": TABLES / "table_s5_permutation_falsification.csv",
    "TABLE_S6": TABLES / "table_s6_accounting.csv",
}

MAIN_TABLE_CAPTIONS = {
    "TABLE_1": "Table 1. Study locations and rainfall observations.",
    "TABLE_2": "Table 2. Observed-weather pseudo-event results.",
    "TABLE_3": "Table 3. Monte Carlo model decomposition.",
    "TABLE_4": "Table 4. Analytic proposition classification.",
}

SUPPLEMENT_TABLE_CAPTIONS = {
    "TABLE_S1": "Table S1. Scenario price models.",
    "TABLE_S2": "Table S2. Rainfall-threshold and attribution-window sensitivity.",
    "TABLE_S3": "Table S3. Simulation parameter sensitivity.",
    "TABLE_S4": "Table S4. Monte Carlo convergence and precision.",
    "TABLE_S5": "Table S5. Permutation falsification.",
    "TABLE_S6": "Table S6. Simulated transfer, cost, profit, and surplus accounting.",
}


def f(value: float, digits: int = 3) -> str:
    return f"{value:.{digits}f}"


def simulation_value(model: str, outcome: str, field: str = "mean") -> float:
    rows = [
        row
        for row in VALUES["simulation"]["summary"]
        if row["model"] == model and row["outcome"] == outcome
    ]
    if len(rows) != 1:
        raise ValueError(f"Missing simulation value: {model} {outcome}")
    return float(rows[0][field])


def pseudo_value(rule: str, field: str) -> float:
    rows = [
        row
        for row in VALUES["empirical"]["pseudo_events"]
        if row["Selection rule"] == rule
    ]
    if len(rows) != 1:
        raise ValueError(f"Missing pseudo-event value: {rule} {field}")
    return float(rows[0][field])


def model_row(model: str) -> dict:
    rows = [
        row
        for row in VALUES["empirical"]["price_models"]
        if row["model"] == model and row["term"] == "scenario_wtp"
    ]
    if len(rows) != 1:
        raise ValueError(f"Missing price model: {model}")
    return rows[0]


def cluster_row() -> dict:
    rows = [
        row
        for row in VALUES["empirical"]["cluster_robust_price_model"]
        if row["term"] == "scenario_wtp"
    ]
    if len(rows) != 1:
        raise ValueError("Missing clustered scenario WTP estimate")
    return rows[0]


def replacements() -> dict[str, str]:
    naive = model_row("naive")
    adjusted = model_row("hazard-adjusted")
    cluster = cluster_row()
    permutation = VALUES["empirical"]["permutation_falsification"][0]
    max_mcse = max(
        float(row["monte_carlo_se"])
        for row in VALUES["simulation"]["convergence"]
        if row["sample_size"] == 10000 and row["outcome"] == "apparent_success"
    )
    rules = {
        "DEFICIT": "30-day deficit threshold",
        "DRYSPELL": "dry-spell threshold",
        "HAZARD": "high prospective hazard",
        "RANDOM": "random drought day",
        "PLACEBO": "season-matched placebo",
    }
    output = {
        "TOTAL_REGION_DAYS": str(VALUES["study"]["region_days"]),
        "TOTAL_PSEUDO_EVENTS": str(
            sum(int(row["Events"]) for row in VALUES["empirical"]["pseudo_events"])
        ),
        "SIM_REPLICATIONS": f"{VALUES['simulation']['replications_per_model']:,}",
        "NAIVE_WTP_COEF": f(float(naive["coefficient"])),
        "NAIVE_CI_LOW": f(float(naive["ci_low"])),
        "NAIVE_CI_HIGH": f(float(naive["ci_high"])),
        "ADJUSTED_WTP_COEF": f(float(adjusted["coefficient"])),
        "ADJUSTED_CI_LOW": f(float(adjusted["ci_low"])),
        "ADJUSTED_CI_HIGH": f(float(adjusted["ci_high"])),
        "CLUSTER_CI_LOW": f(float(cluster["ci_low"])),
        "CLUSTER_CI_HIGH": f(float(cluster["ci_high"])),
        "CLUSTERS": str(int(cluster["clusters"])),
        "RULE_CLUSTERS": str(int(cluster["region_rule_clusters"])),
        "REGRESSION_N": str(int(cluster["n"])),
        "PERM_OBSERVED": f(float(permutation["observed_adjusted_ols_coefficient"]), 4),
        "PERM_MEAN": f(float(permutation["permutation_mean"]), 4),
        "PERMUTATION_P": f(float(permutation["permutation_p_two_sided"])),
        "M0_EVENTS": f(simulation_value("M0_random_timing", "events"), 2),
        "M1_EVENTS": f(simulation_value("M1_drought_demand", "events"), 2),
        "M3_EVENTS": f(simulation_value("M3_strategic_selection", "events"), 2),
        "M0_SUCCESS": f(simulation_value("M0_random_timing", "apparent_success")),
        "M1_SUCCESS": f(simulation_value("M1_drought_demand", "apparent_success")),
        "M2_SUCCESS": f(simulation_value("M2_realistic_hazard", "apparent_success")),
        "M3_SUCCESS": f(
            simulation_value("M3_strategic_selection", "apparent_success")
        ),
        "M4_PRICE": f(simulation_value("M4_bayesian_reputation", "mean_price")),
        "M5_PRICE": f(simulation_value("M5_selective_recency", "mean_price")),
        "FIXED_PRICE": f(float(VALUES["simulation"]["parameters"]["fixed_price"])),
        "M4_PRICE_PREMIUM": f(
            simulation_value("M4_bayesian_reputation", "mean_price")
            - float(VALUES["simulation"]["parameters"]["fixed_price"])
        ),
        "M5_PRICE_PREMIUM": f(
            simulation_value("M5_selective_recency", "mean_price")
            - float(VALUES["simulation"]["parameters"]["fixed_price"])
        ),
        "M4_PROVIDER_PRICE_DISPERSION": f(
            simulation_value(
                "M4_bayesian_reputation", "provider_price_dispersion"
            )
        ),
        "M5_PROVIDER_PRICE_DISPERSION": f(
            simulation_value(
                "M5_selective_recency", "provider_price_dispersion"
            )
        ),
        "M4_REPUTATION": f(
            simulation_value("M4_bayesian_reputation", "reputation_dispersion")
        ),
        "M5_REPUTATION": f(
            simulation_value("M5_selective_recency", "reputation_dispersion")
        ),
        "M4_PRICE_SUCCESS": f(
            simulation_value(
                "M4_bayesian_reputation", "price_success_difference"
            )
        ),
        "M5_PRICE_SUCCESS": f(
            simulation_value("M5_selective_recency", "price_success_difference")
        ),
        "MAX_SUCCESS_MCSE": f(max_mcse, 4),
    }
    for token, rule in rules.items():
        output[f"{token}_SUCCESS_7D"] = f(
            pseudo_value(rule, "Apparent_success_7d")
        )
    return output


def crossref_message(doi: str) -> dict:
    slug = doi.lower().replace("/", "_").replace(".", "-")
    path = ROOT / "data" / "raw" / "literature" / f"crossref_{slug}.json"
    return json.loads(path.read_text())["message"]


def author_string(message: dict) -> str:
    authors = message.get("author", [])
    formatted = []
    for author in authors[:6]:
        family = author.get("family", "").strip()
        if family.isupper():
            family = family.title()
        initials = "".join(
            part[0].upper()
            for part in re.findall(r"[A-Za-zÀ-ÿ]+", author.get("given", ""))
        )
        formatted.append(f"{family} {initials}".strip())
    if len(authors) > 6:
        formatted.append("et al")
    return ", ".join(formatted)


def format_reference(doi: str) -> str:
    if doi == "10.14989/68028":
        return (
            "Akong'a J. Rainmaking rituals: a comparative study of two Kenyan "
            "societies. African Study Monographs. 1987;8(2):71–85. "
            "doi:10.14989/68028."
        )
    message = crossref_message(doi)
    title = message.get("title", [""])[0].rstrip(".")
    container = message.get("container-title", [""])[0]
    issued = message.get("published-print") or message.get("published-online") or message.get("issued")
    year = issued["date-parts"][0][0]
    volume = message.get("volume", "")
    issue = message.get("issue", "")
    pages = message.get("page", "")
    bibliographic = str(year)
    if volume:
        bibliographic += f";{volume}"
        if issue:
            bibliographic += f"({issue})"
        if pages:
            bibliographic += f":{pages}"
    elif pages:
        bibliographic += f":{pages}"
    return (
        f"{author_string(message)}. {title}. {container}. {bibliographic}. "
        f"doi:{doi}."
    )


def reference_block() -> str:
    return "\n".join(
        f"{index}. {format_reference(doi)}"
        for index, doi in enumerate(REFERENCE_ORDER, start=1)
    )


def render_template(source: Path, destination: Path) -> str:
    text = source.read_text()
    for token, value in replacements().items():
        text = text.replace(f"{{{{{token}}}}}", value)
    text = text.replace("{{REFERENCES}}", reference_block())
    unresolved = re.findall(r"\{\{[A-Z0-9_]+\}\}", text)
    allowed = set(FIGURE_FILES) | set(TABLE_FILES)
    unexpected = [token for token in unresolved if token[2:-2] not in allowed]
    if unexpected:
        raise ValueError(f"Unresolved manuscript tokens: {unexpected}")
    destination.write_text(text)
    return text


def add_inline_runs(paragraph, text: str) -> None:
    text = text.replace(r"\(", "").replace(r"\)", "")
    replacements = {
        r"\beta": "β",
        r"\in": "∈",
        r"\mathbf{1}": "1",
        r"\max": "max",
        r"\geq": "≥",
        r"\ldots": "…",
        r"\left": "",
        r"\right": "",
        r"\{": "{",
        r"\}": "}",
    }
    for source, target in replacements.items():
        text = text.replace(source, target)
    text = re.sub(r"_\{([^}]+)\}", r"_\1", text)
    parts = re.split(r"(\*\*.*?\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        else:
            paragraph.add_run(part)


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def add_table(document: Document, marker: str) -> None:
    frame = pd.read_csv(TABLE_FILES[marker])
    caption_text = MAIN_TABLE_CAPTIONS.get(
        marker, SUPPLEMENT_TABLE_CAPTIONS.get(marker, marker.replace("_", " "))
    )
    caption = document.add_paragraph()
    caption.paragraph_format.space_before = Pt(16)
    caption.paragraph_format.space_after = Pt(6)
    caption_run = caption.add_run(caption_text)
    caption_run.bold = True
    table = document.add_table(rows=1, cols=len(frame.columns))
    table.style = "Table Grid"
    table.autofit = True
    header = table.rows[0]
    header._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for cell, column in zip(header.cells, frame.columns):
        cell.text = str(column).replace("_", " ")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, "D9EAF7")
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.size = Pt(7)
    for row in frame.itertuples(index=False, name=None):
        table_row = table.add_row()
        row_properties = table_row._tr.get_or_add_trPr()
        row_properties.append(OxmlElement("w:cantSplit"))
        cells = table_row.cells
        for cell, value in zip(cells, row):
            if isinstance(value, float):
                cell.text = f"{value:.3f}"
            else:
                cell.text = str(value)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(7)


def add_figure(document: Document, marker: str) -> None:
    number = int(marker.split("_")[1])
    document.add_picture(str(FIGURE_FILES[marker]), width=Inches(6.25))
    image_paragraph = document.paragraphs[-1]
    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption = document.add_paragraph()
    caption.paragraph_format.space_before = Pt(12)
    caption.paragraph_format.space_after = Pt(8)
    run = caption.add_run(
        f"Figure {number}. {VALUES['figure_captions'][f'Figure {number}']}"
    )
    run.bold = True


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    normal.paragraph_format.space_after = Pt(0)
    for name in ["Title", "Heading 1", "Heading 2", "Heading 3"]:
        styles[name].font.name = "Arial"
    if "Figure Caption" not in styles:
        style = styles.add_style("Figure Caption", WD_STYLE_TYPE.PARAGRAPH)
        style.font.name = "Times New Roman"
        style.font.size = Pt(10)


def markdown_to_docx(markdown: str, destination: Path) -> None:
    document = Document()
    configure_document(document)
    blocks = re.split(r"\n\s*\n", markdown.strip())
    for block in blocks:
        lines = [line.strip() for line in block.splitlines() if line.strip()]
        if not lines:
            continue
        if lines[0].startswith("- "):
            items = []
            current = ""
            for line in lines:
                if line.startswith("- "):
                    if current:
                        items.append(current)
                    current = line[2:]
                else:
                    current += " " + line
            if current:
                items.append(current)
            for item in items:
                paragraph = document.add_paragraph(style="List Bullet")
                add_inline_runs(paragraph, item)
            continue
        if re.match(r"^\d+\. ", lines[0]):
            items = []
            current = ""
            for line in lines:
                if re.match(r"^\d+\. ", line):
                    if current:
                        items.append(current)
                    current = line
                else:
                    current += " " + line
            if current:
                items.append(current)
            for item in items:
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.left_indent = Inches(0.25)
                paragraph.paragraph_format.first_line_indent = Inches(-0.25)
                add_inline_runs(paragraph, item)
            continue
        if lines[0] == r"\[" and lines[-1] == r"\]":
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_inline_runs(paragraph, " ".join(lines[1:-1]))
            continue
        line = " ".join(lines)
        marker_match = re.fullmatch(r"\{\{([A-Z0-9_]+)\}\}", line)
        if marker_match:
            marker = marker_match.group(1)
            if marker in FIGURE_FILES:
                add_figure(document, marker)
            elif marker in TABLE_FILES:
                add_table(document, marker)
            continue
        if line.startswith("# "):
            paragraph = document.add_paragraph(style="Title")
            add_inline_runs(paragraph, line[2:])
            continue
        if line.startswith("## "):
            paragraph = document.add_heading(level=1)
            add_inline_runs(paragraph, line[3:])
            continue
        if line.startswith("### "):
            paragraph = document.add_heading(level=2)
            add_inline_runs(paragraph, line[4:])
            continue
        paragraph = document.add_paragraph()
        add_inline_runs(paragraph, line)
    document.save(destination)


def manuscript_word_count(text: str) -> int:
    body = text.split("## References", maxsplit=1)[0]
    body = re.sub(r"\{\{[A-Z0-9_]+\}\}", "", body)
    return len(re.findall(r"\b[\w’'-]+\b", body))


def abstract_word_count(text: str) -> int:
    abstract = text.split("## Abstract", 1)[1].split("**Keywords:**", 1)[0]
    return len(re.findall(r"\b[\w’'-]+\b", abstract))


def make_title_page() -> None:
    document = Document()
    configure_document(document)
    title = document.add_paragraph(style="Title")
    title.add_run(
        "Pricing Apparent Environmental Performance: Strategic Selection and "
        "Reputation under a Zero-Effect Rainfall Null"
    )
    for label, value in [
        ("Article type", "Analysis"),
        ("Author(s)", "[AUTHOR INPUT REQUIRED]"),
        ("Affiliation(s)", "[AUTHOR INPUT REQUIRED]"),
        ("Corresponding author", "[AUTHOR INPUT REQUIRED]"),
        ("Email and postal address", "[AUTHOR INPUT REQUIRED]"),
        ("Telephone number", "[AUTHOR INPUT REQUIRED]"),
        ("ORCID(s)", "[AUTHOR INPUT REQUIRED, if applicable]"),
    ]:
        paragraph = document.add_paragraph()
        run = paragraph.add_run(f"{label}: ")
        run.bold = True
        paragraph.add_run(value)
    document.save(MANUSCRIPT / "title_page.docx")


def make_supporting_files(word_count: int, abstract_count: int) -> None:
    total_replications = len(VALUES["simulation"]["models"]) * int(
        VALUES["simulation"]["replications_per_model"]
    )
    location_count = int(VALUES["study"]["locations"])
    highlights = [
        "Natural rainfall creates apparent success under exact zero efficacy",
        "Strategic acceptance changes who is observed, not rainfall",
        "Past attributed outcomes create simulated reputation price premia",
        "A positive current price-success association is not supported",
        "Scenario prices are distinct from causal and plural values",
    ]
    if any(len(item) > 85 for item in highlights):
        raise ValueError("A highlight exceeds the 85-character limit")
    (MANUSCRIPT / "highlights.txt").write_text(
        "\n".join(f"• {item}" for item in highlights) + "\n"
    )
    cover = f"""# Cover letter draft

Dear Editors,

Please consider our Analysis manuscript, “Pricing Apparent Environmental Performance:
Strategic Selection and Reputation under a Zero-Effect Rainfall Null,” for publication
in *Ecological Economics*.

The paper studies an ecological-economic system in which environmental scarcity drives
demand, providers strategically select requests, observers attribute natural rainfall
to accepted events, and valuation and reputation respond—even though the intervention
has exactly zero causal effect on rainfall. Using archived ERA5 rainfall at
{location_count} contrasting locations and {total_replications:,} principal Monte
Carlo replications, we distinguish
natural hazard, apparent efficacy, scenario willingness to pay, transfers, causal
rainfall value, and separate social utility.

Espín-Sánchez, Gil-Guirado, and Ryan establish the timing, hazard, persuasion, and
historical rain-prayer foundation. Our qualified extension separates requests from
provider acceptance, models provider-specific reputation stocks and simulated prices,
and reports the unfavorable principal result rather than tuning it away: reputation
price premia coexist with negative mean contemporaneous price–success differences.
The manuscript addresses the journal's interest in ecological uncertainty, valuation,
behavior, institutions, information asymmetry, and plural values.

[AUTHOR CONFIRMATION REQUIRED: The manuscript is original and is not under consideration
elsewhere.] It uses no human participants or confidential data. All redistributable
analytical inputs, code, acquisition records, checksums, simulation outputs, and
manuscript values accompany the submission; exact retrieval records are supplied for
audit sources that are not redistributed.

[AUTHOR INPUT REQUIRED: corresponding author name, affiliation, contact details,
funding, competing interests, acknowledgements, and any suggested reviewers.]

Sincerely,

[AUTHOR INPUT REQUIRED]
"""
    (MANUSCRIPT / "cover_letter.md").write_text(cover)
    cover_docx = MANUSCRIPT / "cover_letter.docx"
    markdown_to_docx(cover, cover_docx)
    convert_to_pdf(cover_docx)
    declarations = """# Submission declarations

- **Funding:** [AUTHOR INPUT REQUIRED]
- **Competing interests:** [AUTHOR INPUT REQUIRED]
- **CRediT authorship:** [AUTHOR INPUT REQUIRED]
- **Acknowledgements:** [AUTHOR INPUT REQUIRED]
- **Ethics:** Public environmental data, bibliographic metadata, and fictional
  simulations only; no human participants, personal data, or animal research.
- **Data and code:** Public code, legally redistributable inputs, frozen outputs, and
  acquisition ledgers: https://github.com/bougtoir/pricing-the-rain.
- **Generative AI:** During preparation of this work, the authors used Devin
  (Cognition AI) to assist with code development, manuscript drafting, and language
  organization. [AUTHOR CONFIRMATION REQUIRED: after using this tool, the authors
  reviewed and edited the content as needed and take full responsibility for the
  content of the publication.] No generative-AI image tool was used to create or alter
  scientific figures.
"""
    (MANUSCRIPT / "declarations.md").write_text(declarations)
    metrics = {
        "main_text_words_excluding_references": word_count,
        "abstract_words": abstract_count,
        "references": len(REFERENCE_ORDER),
        "figures": len(FIGURE_FILES),
        "main_tables": 4,
        "supplementary_tables": 6,
        "highlights": len(highlights),
        "max_highlight_characters": max(map(len, highlights)),
    }
    (MANUSCRIPT / "manuscript_metrics.json").write_text(
        json.dumps(metrics, indent=2) + "\n"
    )


def make_submission_archive() -> None:
    archive_path = MANUSCRIPT / (
        "pricing_apparent_environmental_performance_submission.zip"
    )
    package_files = [
        MANUSCRIPT / "manuscript_text.docx",
        MANUSCRIPT / "manuscript_text.pdf",
        MANUSCRIPT / "manuscript_blinded.docx",
        MANUSCRIPT / "manuscript_blinded.pdf",
        MANUSCRIPT / "title_page.docx",
        MANUSCRIPT / "highlights.txt",
        MANUSCRIPT / "cover_letter.docx",
        MANUSCRIPT / "cover_letter.pdf",
        MANUSCRIPT / "declarations.md",
        SUPPLEMENT / "supplement.docx",
        SUPPLEMENT / "supplement.pdf",
        FIGURES / "figures_editable.pptx",
        TABLES / "main_tables_editable.docx",
        TABLES / "supplementary_tables_editable.docx",
        ROOT / "AUTHOR_ACTIONS_REQUIRED.md",
        ROOT / "docs" / "JOURNAL_COMPLIANCE_CHECKLIST.md",
    ]
    for stem in [
        "figure1_framework",
        "figure2_empirical_hazard",
        "figure3_pseudo_event_success",
        "figure4_simulation_decomposition",
        "figure5_robustness",
    ]:
        package_files.extend(
            [
                FIGURES / f"{stem}.pdf",
                FIGURES / f"{stem}.png",
            ]
        )
    with zipfile.ZipFile(archive_path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for source in package_files:
            if not source.exists():
                raise FileNotFoundError(f"Missing submission artifact: {source}")
            info = zipfile.ZipInfo(str(source.relative_to(ROOT)))
            info.date_time = (1980, 1, 1, 0, 0, 0)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o100644 << 16
            archive.writestr(info, source.read_bytes())


def make_submission_text(rendered: str) -> str:
    text = rendered
    for marker in FIGURE_FILES:
        text = text.replace(f"{{{{{marker}}}}}", "")
    legends = "\n\n".join(
        f"**Figure {number}.** {VALUES['figure_captions'][f'Figure {number}']}"
        for number in range(1, 6)
    )
    return text.rstrip() + "\n\n## Figure legends\n\n" + legends + "\n"


def convert_to_pdf(path: Path) -> None:
    subprocess.run(
        [
            "libreoffice",
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(path.parent),
            str(path),
        ],
        check=True,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        text=True,
    )


def main() -> None:
    rendered = render_template(
        MANUSCRIPT / "manuscript_template.md", MANUSCRIPT / "manuscript.md"
    )
    supplement = render_template(
        SUPPLEMENT / "supplement_template.md", SUPPLEMENT / "supplement.md"
    )
    main_docx = MANUSCRIPT / "manuscript_blinded.docx"
    submission_docx = MANUSCRIPT / "manuscript_text.docx"
    supplement_docx = SUPPLEMENT / "supplement.docx"
    markdown_to_docx(rendered, main_docx)
    submission_text = make_submission_text(rendered)
    (MANUSCRIPT / "manuscript_text.md").write_text(submission_text)
    markdown_to_docx(submission_text, submission_docx)
    markdown_to_docx(supplement, supplement_docx)
    make_title_page()
    convert_to_pdf(main_docx)
    convert_to_pdf(submission_docx)
    convert_to_pdf(supplement_docx)
    words = manuscript_word_count(rendered)
    abstract_words = abstract_word_count(rendered)
    make_supporting_files(words, abstract_words)
    make_submission_archive()
    print(
        f"Built manuscript ({words} words; abstract {abstract_words} words), "
        "supplement, title page, highlights, cover letter, and submission archive"
    )


if __name__ == "__main__":
    main()
