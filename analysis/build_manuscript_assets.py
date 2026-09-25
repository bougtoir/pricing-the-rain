from __future__ import annotations

import json
from pathlib import Path

import cairosvg
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import seaborn as sns
import yaml
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from PIL import Image
from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches as PptxInches
from pptx.util import Pt as PptxPt


ROOT = Path(__file__).resolve().parents[1]
FIGURES = ROOT / "figures"
TABLES = ROOT / "tables"
RESULTS = ROOT / "results"
ANALYSIS = ROOT / "analysis"
TABLES.mkdir(exist_ok=True)

MODEL_LABELS = {
    "M0_random_timing": "M0 Random timing",
    "M1_drought_demand": "M1 Drought demand",
    "M2_realistic_hazard": "M2 Realistic hazard",
    "M3_strategic_selection": "M3 Strategic selection",
    "M4_bayesian_reputation": "M4 Bayesian reputation",
    "M5_selective_recency": "M5 Selective/recency memory",
}

FIGURE_CAPTIONS = {
    "Figure 1": (
        "Conceptual framework. Environmental hazard enters demand, provider selection, "
        "valuation, attribution, and reputation, while ritual never changes rainfall."
    ),
    "Figure 2": (
        "Estimated seven-day rainfall hazards by preceding dry-spell duration and "
        "location in the 1991–2010 training period. The eight locations are "
        "pre-specified mechanism stress tests, not a representative global sample."
    ),
    "Figure 3": (
        "Natural rainfall following fictional zero-effect pseudo-events in 2011–2020, "
        "shown as apparent success by prospective event-selection rule."
    ),
    "Figure 4": (
        "Sequential Monte Carlo model enrichment under M0–M5. M0–M1 shuffle rainfall "
        "outcomes within sampled years; M2–M5 retain observed local and seasonal "
        "ordering. Strategic selection changes accepted-event composition and current "
        "apparent success. In M4–M5, past attributed outcomes accumulate in reputation "
        "stocks that create price premia and provider-price heterogeneity, while "
        "current success remains a natural-rainfall flow governed by current hazard. "
        "Mean price(success) minus price(failure) is negative in M4–M5. Error bars are "
        "2.5th and 97.5th Monte Carlo percentiles, not confidence intervals for "
        "historical parameters."
    ),
    "Figure 5": (
        "Robustness. Apparent success across rainfall thresholds and windows, selected "
        "parameter sensitivity, and Monte Carlo precision."
    ),
}


def read_config() -> tuple[dict, dict]:
    regions = yaml.safe_load((ROOT / "config" / "regions.yml").read_text())
    simulation = yaml.safe_load((ROOT / "config" / "simulation.yml").read_text())
    return regions, simulation


def save_table(frame: pd.DataFrame, stem: str) -> None:
    frame.to_csv(TABLES / f"{stem}.csv", index=False)


def make_tables(regions_config: dict) -> dict[str, pd.DataFrame]:
    weather = pd.read_csv(
        ROOT / "data" / "processed" / "daily_weather_features.csv.gz",
        usecols=["region", "date", "precip_mm"],
    )
    region_rows = []
    for region in regions_config["regions"]:
        subset = weather.loc[weather["region"] == region["id"]]
        region_rows.append(
            {
                "Location": region["name"],
                "Latitude": region["latitude"],
                "Longitude": region["longitude"],
                "Rainfall regime": region["regime"],
                "Region-days": int(len(subset)),
                "Missing precipitation": int(subset["precip_mm"].isna().sum()),
            }
        )
    table1 = pd.DataFrame(region_rows)

    pseudo = pd.read_csv(ANALYSIS / "pseudo_events.csv")
    table2 = (
        pseudo.groupby("rule", observed=True)
        .agg(
            Events=("rule", "size"),
            Apparent_success_1d=("success_1mm_1d", "mean"),
            Apparent_success_3d=("success_1mm_3d", "mean"),
            Apparent_success_7d=("success_1mm_7d", "mean"),
            Apparent_success_14d=("success_1mm_14d", "mean"),
            Mean_scenario_WTP=("scenario_wtp", "mean"),
        )
        .reset_index()
        .rename(columns={"rule": "Selection rule"})
    )

    summary = pd.read_csv(RESULTS / "simulation_summary.csv")
    selected = [
        "events",
        "apparent_success",
        "mean_price",
        "provider_price_dispersion",
        "price_success_difference",
        "reputation_dispersion",
    ]
    table3 = (
        summary.loc[summary["outcome"].isin(selected)]
        .pivot(index="model", columns="outcome", values="mean")
        .reindex(MODEL_LABELS)
        .reset_index()
        .rename(columns={"model": "Model"})
    )
    table3["Model"] = table3["Model"].map(MODEL_LABELS)

    table4 = pd.read_csv(RESULTS / "hypothesis_classification.csv").rename(
        columns={
            "hypothesis": "Hypothesis",
            "classification": "Classification",
            "basis": "Basis",
        }
    )

    regression = pd.read_csv(ANALYSIS / "price_success_models.csv")
    cluster = pd.read_csv(ANALYSIS / "cluster_robust_price_success.csv")
    permutation = pd.read_csv(ANALYSIS / "price_permutation_falsification.csv")
    table_s1 = pd.concat(
        [
            regression.assign(inference="model-based"),
            cluster.rename(
                columns={"cluster_standard_error": "standard_error"}
            ).assign(
                model="hazard-adjusted",
                inference="two-way clustered: region-year and region-rule",
            ),
        ],
        ignore_index=True,
        sort=False,
    )
    table_s2 = pd.read_csv(ANALYSIS / "threshold_window_sensitivity.csv")
    table_s3 = pd.read_csv(RESULTS / "simulation_sensitivity_summary.csv")
    table_s4 = pd.read_csv(RESULTS / "simulation_convergence.csv")
    table_s5 = permutation
    accounting = [
        "provider_revenue",
        "community_expenditure",
        "provider_resource_cost",
        "provider_profit",
        "community_material_social_utility",
        "total_social_surplus",
        "provider_revenue_dispersion",
    ]
    table_s6 = (
        summary.loc[summary["outcome"].isin(accounting)]
        .pivot(index="model", columns="outcome", values="mean")
        .reindex(MODEL_LABELS)
        .reset_index()
        .rename(columns={"model": "Model"})
    )
    table_s6["Model"] = table_s6["Model"].map(MODEL_LABELS)

    tables = {
        "table1_study_locations": table1,
        "table2_pseudo_event_results": table2,
        "table3_simulation_results": table3,
        "table4_hypothesis_classification": table4,
        "table_s1_price_models": table_s1,
        "table_s2_threshold_window_sensitivity": table_s2,
        "table_s3_parameter_sensitivity": table_s3,
        "table_s4_simulation_convergence": table_s4,
        "table_s5_permutation_falsification": table_s5,
        "table_s6_accounting": table_s6,
    }
    for stem, frame in tables.items():
        save_table(frame, stem)
    return tables


def make_figure5() -> None:
    sensitivity = pd.read_csv(ANALYSIS / "threshold_window_sensitivity.csv")
    sensitivity[["threshold", "window"]] = sensitivity["outcome"].str.extract(
        r"success_([0-9.]+)mm_([0-9]+)d"
    )
    sensitivity["window"] = sensitivity["window"].astype(int)
    heat = (
        sensitivity.groupby(["threshold", "window"], observed=True)[
            "apparent_success"
        ]
        .mean()
        .unstack()
        .reindex(["0.1", "1", "5"])
    )

    parameter = pd.read_csv(RESULTS / "simulation_sensitivity_summary.csv")
    parameter = parameter.loc[
        ((parameter["scenario"].str.startswith("M3_")) & (parameter["outcome"] == "apparent_success"))
        | (
            parameter["scenario"].str.startswith(("M4_", "M5_"))
            & (parameter["outcome"] == "mean_price")
        )
    ].copy()
    parameter["label"] = (
        parameter["scenario"]
        .str.replace("M3_selection_slope_", "M3 slope ", regex=False)
        .str.replace("M4_prior_mean_", "M4 prior ", regex=False)
        .str.replace("M5_", "M5 ", regex=False)
        .str.replace("_memory", "", regex=False)
        .str.replace("_", " ", regex=False)
    )

    convergence = pd.read_csv(RESULTS / "simulation_convergence.csv")
    convergence = convergence.loc[
        (convergence["outcome"] == "apparent_success")
        & convergence["model"].isin(
            ["M0_random_timing", "M3_strategic_selection", "M5_selective_recency"]
        )
    ].copy()
    convergence["model"] = convergence["model"].map(MODEL_LABELS)

    sns.set_theme(style="whitegrid", context="paper")
    figure, axes = plt.subplots(1, 3, figsize=(13.2, 4.4))
    sns.heatmap(
        heat,
        annot=True,
        fmt=".2f",
        cmap="Blues",
        vmin=0,
        vmax=1,
        cbar_kws={"label": "Apparent success"},
        ax=axes[0],
    )
    axes[0].set_title("A. Threshold and window sensitivity")
    axes[0].set_xlabel("Post-event window (days)")
    axes[0].set_ylabel("Rainfall threshold (mm)")

    sns.barplot(data=parameter, x="mean", y="label", hue="outcome", ax=axes[1])
    axes[1].set_title("B. Parameter sensitivity")
    axes[1].set_xlabel("Mean apparent success or price")
    axes[1].set_ylabel("")
    axes[1].legend(title="", loc="lower right", fontsize=7)

    sns.lineplot(
        data=convergence,
        x="sample_size",
        y="monte_carlo_se",
        hue="model",
        marker="o",
        ax=axes[2],
    )
    axes[2].set_title("C. Monte Carlo precision")
    axes[2].set_xlabel("Replications")
    axes[2].set_ylabel("Monte Carlo standard error")
    axes[2].legend(title="", fontsize=7)
    figure.tight_layout()
    figure.savefig(FIGURES / "figure5_robustness.pdf", bbox_inches="tight")
    figure.savefig(
        FIGURES / "figure5_robustness.png", dpi=600, bbox_inches="tight"
    )
    plt.close(figure)


def add_docx_table(document: Document, title: str, frame: pd.DataFrame) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(16)
    run = paragraph.add_run(title)
    run.bold = True
    table = document.add_table(rows=1, cols=len(frame.columns))
    table.style = "Table Grid"
    for cell, column in zip(table.rows[0].cells, frame.columns):
        cell.text = str(column).replace("_", " ")
    for row in frame.itertuples(index=False, name=None):
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            if isinstance(value, float):
                cell.text = f"{value:.3f}"
            else:
                cell.text = str(value)


def make_tables_docx(tables: dict[str, pd.DataFrame]) -> None:
    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    title = document.add_heading(
        "Pricing Apparent Environmental Performance — Editable Tables", level=0
    )
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    main_titles = {
        "table1_study_locations": "Table 1. Study locations and rainfall observations",
        "table2_pseudo_event_results": "Table 2. Observed-weather pseudo-event results",
        "table3_simulation_results": "Table 3. Monte Carlo model decomposition",
        "table4_hypothesis_classification": "Table 4. Analytic proposition classification",
    }
    for stem, title_text in main_titles.items():
        add_docx_table(document, title_text, tables[stem])
        document.add_page_break()
    document.save(TABLES / "main_tables_editable.docx")

    supplement = Document()
    supplement.sections[0].orientation = WD_ORIENT.LANDSCAPE
    supplement.sections[0].page_width, supplement.sections[0].page_height = (
        supplement.sections[0].page_height,
        supplement.sections[0].page_width,
    )
    supplement.add_heading(
        "Pricing Apparent Environmental Performance — Editable Supplementary Tables",
        level=0,
    )
    for stem in [
        "table_s1_price_models",
        "table_s2_threshold_window_sensitivity",
        "table_s3_parameter_sensitivity",
        "table_s4_simulation_convergence",
        "table_s5_permutation_falsification",
        "table_s6_accounting",
    ]:
        add_docx_table(
            supplement,
            stem.replace("_", " ").replace("table s", "Table S").title(),
            tables[stem],
        )
        supplement.add_page_break()
    supplement.save(TABLES / "supplementary_tables_editable.docx")


def make_figures_pptx() -> None:
    png1 = FIGURES / "figure1_framework.png"
    cairosvg.svg2pdf(
        url=str(FIGURES / "figure1_framework.svg"),
        write_to=str(FIGURES / "figure1_framework.pdf"),
    )
    cairosvg.svg2png(
        url=str(FIGURES / "figure1_framework.svg"),
        write_to=str(png1),
        output_width=2400,
    )
    figure_paths = [
        png1,
        FIGURES / "figure2_empirical_hazard.png",
        FIGURES / "figure3_pseudo_event_success.png",
        FIGURES / "figure4_simulation_decomposition.png",
        FIGURES / "figure5_robustness.png",
    ]
    presentation = Presentation()
    presentation.slide_width = PptxInches(13.333)
    presentation.slide_height = PptxInches(7.5)
    for index, path in enumerate(figure_paths, start=1):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        title_box = slide.shapes.add_textbox(
            PptxInches(0.5), PptxInches(0.15), PptxInches(12.3), PptxInches(0.45)
        )
        title = title_box.text_frame.paragraphs[0]
        title.text = f"Figure {index}"
        title.font.size = PptxPt(22)
        title.font.bold = True
        title.alignment = PP_ALIGN.CENTER
        image_width, image_height = Image.open(path).size
        image_ratio = image_width / image_height
        box_width, box_height = 11.93, 5.65
        if image_ratio >= box_width / box_height:
            width = box_width
            height = width / image_ratio
        else:
            height = box_height
            width = height * image_ratio
        left = 0.7 + (box_width - width) / 2
        top = 0.7 + (box_height - height) / 2
        slide.shapes.add_picture(
            str(path),
            PptxInches(left),
            PptxInches(top),
            width=PptxInches(width),
            height=PptxInches(height),
        )
        caption_box = slide.shapes.add_textbox(
            PptxInches(0.65), PptxInches(6.45), PptxInches(12.03), PptxInches(0.8)
        )
        caption = caption_box.text_frame.paragraphs[0]
        caption.text = FIGURE_CAPTIONS[f"Figure {index}"]
        caption.font.size = PptxPt(11)
    presentation.save(FIGURES / "figures_editable.pptx")


def json_records(frame: pd.DataFrame) -> list[dict[str, object]]:
    clean = frame.replace({np.nan: None})
    return clean.to_dict(orient="records")


def make_values(tables: dict[str, pd.DataFrame], regions: dict, simulation: dict) -> None:
    payload = {
        "study": {
            "start": regions["study_period"]["start"],
            "end": regions["study_period"]["end"],
            "training_end": regions["training_end"],
            "rain_threshold_mm": regions["primary_rain_threshold_mm"],
            "success_windows_days": regions["success_windows_days"],
            "locations": len(regions["regions"]),
            "region_days": int(tables["table1_study_locations"]["Region-days"].sum()),
        },
        "simulation": {
            "seed": simulation["seed"],
            "replications_per_model": simulation["replications_per_model"],
            "providers": simulation["providers"],
            "models": simulation["models"],
            "parameters": simulation["parameters"],
            "summary": json_records(
                pd.read_csv(RESULTS / "simulation_summary.csv")
            ),
            "sensitivity": json_records(
                pd.read_csv(RESULTS / "simulation_sensitivity_summary.csv")
            ),
            "convergence": json_records(
                pd.read_csv(RESULTS / "simulation_convergence.csv")
            ),
        },
        "empirical": {
            "pseudo_events": json_records(tables["table2_pseudo_event_results"]),
            "price_models": json_records(
                pd.read_csv(ANALYSIS / "price_success_models.csv")
            ),
            "cluster_robust_price_model": json_records(
                pd.read_csv(ANALYSIS / "cluster_robust_price_success.csv")
            ),
            "permutation_falsification": json_records(
                pd.read_csv(ANALYSIS / "price_permutation_falsification.csv")
            ),
            "threshold_window_sensitivity": json_records(
                pd.read_csv(ANALYSIS / "threshold_window_sensitivity.csv")
            ),
        },
        "hypotheses": json_records(
            pd.read_csv(RESULTS / "hypothesis_classification.csv")
        ),
        "figure_captions": FIGURE_CAPTIONS,
    }
    (RESULTS / "manuscript_values.json").write_text(
        json.dumps(payload, indent=2, ensure_ascii=False) + "\n"
    )


def main() -> None:
    regions, simulation = read_config()
    tables = make_tables(regions)
    make_figure5()
    make_tables_docx(tables)
    make_figures_pptx()
    make_values(tables, regions, simulation)
    print("Built manuscript figures, tables, and machine-readable values")


if __name__ == "__main__":
    main()
